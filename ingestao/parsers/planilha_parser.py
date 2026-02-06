"""
Parser para planilhas (XLSX, XLS, CSV).
Comum em relatórios de faturamento, bases de procedimentos,
tabelas de preços e exportações de sistemas hospitalares.
"""
import csv
import io
from typing import Optional

from ingestao.parsers.base import BaseParser, ResultadoParsing


class PlanilhaParser(BaseParser):
    """Parser para planilhas Excel e CSV."""

    def tipos_suportados(self) -> list[str]:
        return ["xlsx", "xls", "csv"]

    def parse(self, conteudo: bytes, nome_arquivo: str) -> ResultadoParsing:
        resultado = ResultadoParsing()
        resultado.tipo_documento = "planilha"

        ext = nome_arquivo.rsplit(".", 1)[-1].lower() if "." in nome_arquivo else ""

        try:
            if ext == "csv":
                self._parse_csv(conteudo, resultado)
            else:
                self._parse_excel(conteudo, resultado)
        except Exception as e:
            resultado.erros.append(f"Erro ao processar planilha: {str(e)}")

        return resultado

    def _parse_csv(self, conteudo: bytes, resultado: ResultadoParsing):
        """Parse de arquivo CSV."""
        texto = conteudo.decode("utf-8", errors="replace")
        leitor = csv.DictReader(io.StringIO(texto), delimiter=self._detectar_delimitador(texto))
        linhas = list(leitor)

        if not linhas:
            resultado.avisos.append("Planilha CSV vazia.")
            return

        resultado.metadados["total_linhas"] = len(linhas)
        resultado.metadados["colunas"] = list(linhas[0].keys()) if linhas else []

        tipo_planilha = self._identificar_tipo_planilha(
            resultado.metadados["colunas"]
        )
        resultado.metadados["tipo_planilha_detectado"] = tipo_planilha

        for linha in linhas:
            self._processar_linha(linha, tipo_planilha, resultado)

    def _parse_excel(self, conteudo: bytes, resultado: ResultadoParsing):
        """Parse de arquivo Excel."""
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True)

            for nome_aba in wb.sheetnames:
                ws = wb[nome_aba]
                linhas = list(ws.iter_rows(values_only=True))
                if not linhas:
                    continue

                cabecalho = [str(c).strip() if c else "" for c in linhas[0]]
                resultado.metadados[f"aba_{nome_aba}_colunas"] = cabecalho
                resultado.metadados[f"aba_{nome_aba}_linhas"] = len(linhas) - 1

                tipo_planilha = self._identificar_tipo_planilha(cabecalho)

                for linha_vals in linhas[1:]:
                    linha_dict = {}
                    for i, col in enumerate(cabecalho):
                        if i < len(linha_vals) and col:
                            linha_dict[col] = linha_vals[i]
                    self._processar_linha(linha_dict, tipo_planilha, resultado)

            wb.close()

        except ImportError:
            resultado.erros.append(
                "Biblioteca openpyxl não instalada. Execute: pip install openpyxl"
            )

    def _detectar_delimitador(self, texto: str) -> str:
        """Detecta o delimitador do CSV."""
        primeira_linha = texto.split("\n")[0]
        contagem = {
            ";": primeira_linha.count(";"),
            ",": primeira_linha.count(","),
            "\t": primeira_linha.count("\t"),
            "|": primeira_linha.count("|"),
        }
        return max(contagem, key=contagem.get)

    def _identificar_tipo_planilha(self, colunas: list[str]) -> str:
        """Identifica o tipo de dados na planilha com base nas colunas."""
        cols_lower = [c.lower() for c in colunas if c]

        # Planilha de procedimentos/faturamento
        if any(
            t in " ".join(cols_lower)
            for t in ["tuss", "procedimento", "código proc", "cod proc"]
        ):
            return "procedimentos"

        # Planilha de glosas
        if any(
            t in " ".join(cols_lower)
            for t in ["glosa", "motivo glosa", "valor glosado"]
        ):
            return "glosas"

        # Planilha de prestadores
        if any(t in " ".join(cols_lower) for t in ["cnes", "prestador"]):
            return "prestadores"

        # Planilha de beneficiários
        if any(
            t in " ".join(cols_lower)
            for t in ["carteirinha", "beneficiário", "carteira"]
        ):
            return "beneficiarios"

        return "generico"

    def _processar_linha(
        self,
        linha: dict,
        tipo_planilha: str,
        resultado: ResultadoParsing,
    ):
        """Processa uma linha da planilha conforme o tipo detectado."""
        if tipo_planilha == "procedimentos":
            proc = self._mapear_procedimento(linha)
            if proc:
                resultado.procedimentos.append(proc)
        elif tipo_planilha == "glosas":
            glosa = self._mapear_glosa(linha)
            if glosa:
                resultado.glosas.append(glosa)
        elif tipo_planilha == "prestadores":
            prest = self._mapear_prestador(linha)
            if prest:
                resultado.prestadores.append(prest)
        elif tipo_planilha == "beneficiarios":
            benef = self._mapear_beneficiario(linha)
            if benef:
                resultado.beneficiarios.append(benef)
        else:
            # Tentar extrair o máximo de dados genéricos
            self._mapear_generico(linha, resultado)

    def _encontrar_valor(self, linha: dict, termos: list[str]) -> Optional[str]:
        """Busca um valor em colunas que contenham os termos dados."""
        for chave, valor in linha.items():
            chave_lower = str(chave).lower()
            for termo in termos:
                if termo in chave_lower and valor is not None:
                    return str(valor).strip()
        return None

    def _mapear_procedimento(self, linha: dict) -> Optional[dict]:
        codigo = self._encontrar_valor(
            linha, ["tuss", "código", "cod proc", "cod_proc", "codigo"]
        )
        if not codigo:
            return None
        return {
            "codigo_tuss": codigo,
            "descricao": self._encontrar_valor(
                linha, ["descrição", "descricao", "procedimento", "nome"]
            ),
            "quantidade_realizada": self._encontrar_valor(
                linha, ["qtd", "quantidade", "quant"]
            ),
            "valor_unitario": self._encontrar_valor(
                linha, ["valor unit", "vl unit", "unitário", "unitario"]
            ),
            "valor_total": self._encontrar_valor(
                linha, ["valor total", "vl total", "total"]
            ),
        }

    def _mapear_glosa(self, linha: dict) -> Optional[dict]:
        valor = self._encontrar_valor(
            linha, ["valor glosado", "glosa", "vl glosa"]
        )
        motivo = self._encontrar_valor(
            linha, ["motivo", "justificativa", "código motivo"]
        )
        if not valor and not motivo:
            return None
        return {
            "valor_glosado": valor,
            "codigo_motivo": motivo,
            "descricao_motivo": self._encontrar_valor(
                linha, ["descrição motivo", "desc motivo"]
            ),
            "codigo_tuss": self._encontrar_valor(
                linha, ["tuss", "procedimento", "código"]
            ),
            "numero_guia": self._encontrar_valor(
                linha, ["guia", "nº guia", "numero guia"]
            ),
        }

    def _mapear_prestador(self, linha: dict) -> Optional[dict]:
        cnes = self._encontrar_valor(linha, ["cnes"])
        if not cnes:
            return None
        return {
            "cnes": cnes,
            "cnpj": self._encontrar_valor(linha, ["cnpj"]),
            "razao_social": self._encontrar_valor(
                linha, ["razão social", "razao social", "nome", "prestador"]
            ),
        }

    def _mapear_beneficiario(self, linha: dict) -> Optional[dict]:
        carteirinha = self._encontrar_valor(
            linha, ["carteirinha", "carteira", "número carteira"]
        )
        if not carteirinha:
            return None
        return {
            "numero_carteirinha": carteirinha,
            "nome": self._encontrar_valor(linha, ["nome", "beneficiário"]),
        }

    def _mapear_generico(self, linha: dict, resultado: ResultadoParsing):
        """Tenta extrair dados de uma linha genérica."""
        proc = self._mapear_procedimento(linha)
        if proc:
            resultado.procedimentos.append(proc)
            return
        glosa = self._mapear_glosa(linha)
        if glosa:
            resultado.glosas.append(glosa)
